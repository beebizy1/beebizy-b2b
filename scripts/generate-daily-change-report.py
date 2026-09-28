from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Beebizy-Daily-Change-Report-2026-09-27.docx"
LOGO = ROOT / "public" / "beebizy-logo.png"

NAVY = "111827"
MUTED = "667085"
YELLOW = "FFD600"
PALE_YELLOW = "FFF8CC"
LIGHT = "F7F8FA"
LINE = "E5E7EB"
GREEN = "16794B"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        props.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=LINE, size="6") -> None:
    props = cell._tc.get_or_add_tcPr()
    borders = props.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def set_repeat_table_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    props.append(repeat)


def prevent_row_split(row) -> None:
    props = row._tr.get_or_add_trPr()
    keep = OxmlElement("w:cantSplit")
    props.append(keep)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)
    paragraph.paragraph_format.space_after = Pt(3)


def add_number(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Number")
    paragraph.add_run(text)
    paragraph.paragraph_format.space_after = Pt(4)


def add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_heading(text, level=level)
    keep_with_next(paragraph)


def add_table(document: Document, headers: list[str], rows: Iterable[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, top=75, bottom=75)
        set_cell_border(cell, NAVY)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(label)
        run.bold = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(WHITE)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else LIGHT)
            set_cell_margins(cell, top=75, bottom=75)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(NAVY)
            if widths:
                cell.width = Inches(widths[index])


def add_callout(document: Document, title: str, text: str, fill=PALE_YELLOW) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    prevent_row_split(table.rows[0])
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    set_cell_border(cell, YELLOW, "10")
    title_p = cell.paragraphs[0]
    title_p.paragraph_format.space_after = Pt(3)
    title_run = title_p.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(10)
    title_run.font.color.rgb = RGBColor.from_string(NAVY)
    body = cell.add_paragraph(text)
    body.paragraph_format.space_after = Pt(0)
    body_run = body.runs[0]
    body_run.font.size = Pt(9.5)
    body_run.font.color.rgb = RGBColor.from_string(NAVY)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.04

    title = document.styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)

    for style_name, size, color in (
        ("Heading 1", 17, NAVY),
        ("Heading 2", 12.5, NAVY),
    ):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)

    for list_name in ("List Bullet", "List Number"):
        style = document.styles[list_name]
        style.font.name = "Aptos"
        style.font.size = Pt(9.2)
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.left_indent = Inches(0.22)
        style.paragraph_format.first_line_indent = Inches(-0.16)


def configure_page(document: Document) -> None:
    settings = document.settings.element
    if settings.find(qn("w:evenAndOddHeaders")) is None:
        settings.append(OxmlElement("w:evenAndOddHeaders"))
    section = document.sections[0]
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.22)

    for header in (section.header, section.even_page_header):
        table = header.add_table(rows=1, cols=2, width=Inches(7.1))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        left, right = table.rows[0].cells
        left.width = Inches(5.8)
        right.width = Inches(1.3)
        left.text = "BEEBIZY  |  CHANGE REPORT"
        left.paragraphs[0].runs[0].font.size = Pt(7.5)
        left.paragraphs[0].runs[0].font.bold = True
        left.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(MUTED)
        if LOGO.exists():
            right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            right.paragraphs[0].add_run().add_picture(str(LOGO), width=Inches(0.68))

    for footer in (section.footer, section.even_page_footer):
        add_page_number(footer.paragraphs[0])


def build() -> None:
    document = Document()
    configure_styles(document)
    configure_page(document)

    label = document.add_paragraph()
    label.paragraph_format.space_before = Pt(24)
    label.paragraph_format.space_after = Pt(7)
    run = label.add_run("DAILY PRODUCT UPDATE")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GREEN)

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(8)
    title.add_run("Beebizy Change Report")

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("Changes completed September 27, 2026")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    meta = document.add_table(rows=3, cols=2)
    meta.style = "Table Grid"
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row, (key, value) in zip(
        meta.rows,
        (
            ("Prepared for", "Team review on September 29, 2026"),
            ("Product", "Beebizy Studio"),
            ("Live preview", "beebizy-studio-preview.vercel.app"),
        ),
    ):
        row.cells[0].width = Inches(1.3)
        row.cells[1].width = Inches(5.8)
        for cell in row.cells:
            set_cell_margins(cell, top=90, bottom=90)
            set_cell_border(cell, LINE)
        set_cell_shading(row.cells[0], LIGHT)
        row.cells[0].paragraphs[0].add_run(key).bold = True
        row.cells[1].paragraphs[0].add_run(value)
        for cell in row.cells:
            for text_run in cell.paragraphs[0].runs:
                text_run.font.size = Pt(9)
                text_run.font.color.rgb = RGBColor.from_string(NAVY)

    document.add_paragraph()
    add_callout(
        document,
        "Meeting overview",
        "Yesterday's work removed the main points of friction reported during testing. The updates make it easier to import existing guest lists, assign responsibilities, manage vendor payments, organize multi-day events, and arrange the workspace around each team's process. All items in this report are complete and available on the preview site.",
    )

    add_heading(document, "At a glance", 1)
    add_table(
        document,
        ["Area", "What changed", "Why it matters"],
        [
            ["Guest imports", "Flexible column matching and attendee counts", "Clients can use existing lists without renaming every column."],
            ["Assignments", "Direct assignment emails and links", "The person responsible can open and complete the exact item."],
            ["Deposits", "Saved records connected to vendors and budgets", "Teams can track contracted, paid, outstanding, and overdue amounts."],
            ["Multi-day events", "One day selector across the event workspace", "Schedules and shifts can be organized by event day."],
            ["Workspace layout", "Event tabs can be rearranged", "Each organization can match Beebizy to its workflow."],
            ["Spreadsheet tools", "Import actions are easier to find", "Existing event information is faster to bring in."],
        ],
        widths=[1.3, 2.55, 3.25],
    )

    document.add_page_break()
    add_heading(document, "Changes completed", 1)

    add_heading(document, "1. More flexible guest-list imports", 2)
    document.add_paragraph("Guest lists no longer need to follow one exact Beebizy template.")
    for item in (
        "Upload a CSV file, paste spreadsheet data, or load a public Google Sheet.",
        "Beebizy suggests which columns contain the guest name, email, organization, guest type, notes, and attendee count.",
        "The user can review and change the suggestions before importing.",
        "Guests can be imported without an email address when the list only contains names.",
        "One row can represent several attendees, such as a guest plus a partner.",
        "Registration, check-in, analytics, and printed lists now count people accurately, not only spreadsheet rows.",
        "Event capacity is checked before the import is completed.",
    ):
        add_bullet(document, item)
    add_callout(
        document,
        "Real-world check",
        "The Mrs. Bench registration sheet was tested successfully. Beebizy read 16 spreadsheet rows representing 32 attendees.",
        fill="EAF7EF",
    )

    add_heading(document, "2. Clearer responsibility emails", 2)
    for item in (
        "Checklist and run-of-show assignments can email the responsible person when they are assigned.",
        "Volunteer messages continue to include the role and shift time.",
        "Each email links directly to the exact assignment, where it can be reviewed, completed, or added to Google Calendar.",
        "Existing verified Beebizy users can be activated without being trapped in a repeated invitation flow.",
        "Organizers can send one combined responsibilities email instead of forwarding several separate messages.",
    ):
        add_bullet(document, item)

    add_heading(document, "3. Reliable deposits and vendor payments", 2)
    for item in (
        "New deposits save correctly.",
        "Vendor commitments and relevant budget items can appear in the payment plan.",
        "Teams can see contracted, paid, outstanding, and overdue totals.",
        "Deposit status can be updated as pending, paid, overdue, or refunded.",
        "Adding a vendor from an event returns the user to the same event and connects the vendor to it.",
    ):
        add_bullet(document, item)

    add_heading(document, "4. Multi-day event organization", 2)
    for item in (
        "Users can add Day 2, Day 3, and additional event days.",
        "The selected day stays active while moving between event sections.",
        "Run-of-show items and volunteer shifts follow the selected day.",
        "Event-wide information, such as the overall budget, remains shared across all days.",
    ):
        add_bullet(document, item)

    add_heading(document, "5. Custom event-tab order", 2)
    for item in (
        "Teams can select Arrange and move event tabs up or down.",
        "The saved order applies across events in the selected Beebizy view.",
        "New sections still appear when a custom order is already saved.",
        "The default order can be restored at any time.",
    ):
        add_bullet(document, item)
    document.add_paragraph("The current preference is saved in the browser being used. A future enhancement can save it to the person's Beebizy account so it follows them to every device.")

    add_heading(document, "6. Easier-to-find spreadsheet imports", 2)
    document.add_paragraph("Import actions are now visible inside the Checklist, Budget, and Floorplan sections. The importer also recognizes more common headings, including action item, completion date, target date, point person, person responsible, task owner, complete, and done.")

    document.add_page_break()
    add_heading(document, "Before and now", 1)
    add_table(
        document,
        ["Before", "Now"],
        [
            ["A guest sheet could fail when its headings did not match Beebizy exactly.", "Beebizy suggests column matches and lets the user correct them before import."],
            ["A list of 16 couples could appear as only 16 registrations.", "Party size is retained, so the system can show 32 people."],
            ["An organizer had to explain where an assigned task lived.", "The assignee can receive a direct link to the exact item."],
            ["Deposits could fail to save or stay disconnected from commitments.", "Deposits save and roll into paid and outstanding totals."],
            ["Adding a vendor could move the user away from the event.", "The vendor is added and the user returns to the event workflow."],
            ["Event tabs appeared in one fixed order.", "Teams can arrange tabs to match their process."],
            ["Multi-day planning was strongest only in the run of show.", "The selected day now follows the user and supports volunteer shifts."],
        ],
        widths=[3.55, 3.55],
    )

    add_heading(document, "Verification and release status", 1)
    add_table(
        document,
        ["Check", "Result"],
        [
            ["Live preview", "Updated and available"],
            ["Automated quality checks", "381 passed"],
            ["Type check", "Passed"],
            ["Code quality check", "Passed with no errors"],
            ["Database updates", "Included for deposits and attendee counts"],
            ["Real spreadsheet test", "Mrs. Bench sheet imported as 16 rows and 32 people"],
        ],
        widths=[2.25, 4.85],
    )
    document.add_paragraph("Four existing development warnings remain. They are internal file-organization notices and do not affect client use of the preview.")

    add_heading(document, "Suggested review for tomorrow's meeting", 1)
    for item in (
        "Import a guest list whose column names do not match Beebizy.",
        "Confirm the total number of people, including party sizes.",
        "Open the registration list, check-in view, and printable guest list.",
        "Assign one checklist item and one run-of-show item, then open the emailed links.",
        "Add a vendor, record a deposit, and update its payment status.",
        "Add a second event day and confirm that the selected day follows the user across tabs.",
        "Rearrange the event tabs and confirm the preferred order.",
    ):
        add_number(document, item)

    add_heading(document, "Decisions to confirm", 1)
    for item in (
        "Should attendee count mean the total people in the party, or the additional guests beyond the named person? Beebizy currently treats it as the total represented by that row.",
        "Should the custom tab order follow each person across all devices? It currently stays with the browser and Beebizy view where it was arranged.",
        "Which two or three real client spreadsheets should be used for the next import test?",
    ):
        add_bullet(document, item)

    add_callout(
        document,
        "Bottom line",
        "Yesterday's release addressed the most important operational feedback from hands-on testing. Beebizy now adapts more easily to the spreadsheets and working habits clients already have, while keeping registrations, assignments, schedules, vendors, and payments connected inside the event.",
    )

    document.core_properties.title = "Beebizy Change Report - September 27, 2026"
    document.core_properties.subject = "Plain-language product change report for the Beebizy team"
    document.core_properties.author = "Beebizy"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
